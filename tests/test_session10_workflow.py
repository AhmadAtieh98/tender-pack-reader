"""Session 10, the AI workflow (tenderpack ai run / resume / submit-batch) end to end on blind rehearsal 02's Addendum
No. 3 added to the real pack, with a RECORDED workflow (tests/fixtures/ai_cassettes/workflow_add03.yaml, hand-written:
no live call is made and none is implied).

The previous evidence build is ingested afresh into a disposable folder (never rehearsals/*/build, never build/). A
run writes only under its staging folder (here a tmp folder) and its work log; the real curation/, config/ and out/
are compared byte for byte before and after, and this process's writes are audited.

Passing these tests shows that the workflow accounts for every provision, checkpoints and resumes, keeps the candidate
apart from the real state and validates what a model returns; it says nothing about how well any real model proposes."""
from __future__ import annotations

import json
import socket
import warnings
from pathlib import Path

import pymupdf
import pytest
import yaml

from ai_fixture import CASSETTES, ROOT, WriteAudit, tree_files
from tenderpack.ai import workflow as W
from tenderpack.ai.candidate import BANNER

PDF = ROOT / "rehearsals/blind-02/input/ADD-03_Addendum_No_3.pdf"
CASSETTE = CASSETTES / "workflow_add03.yaml"
REAL = [ROOT / "curation", ROOT / "config", ROOT / "out"]
quiet = lambda *a, **k: None  # noqa: E731


@pytest.fixture(scope="module")
def prev_build(tmp_path_factory):
    """The preceding state's evidence build (the real pack as received), ingested into a disposable folder."""
    from tenderpack.cli import ingest
    out = tmp_path_factory.mktemp("w10-prev") / "build"
    res = ingest(ROOT / "config/pack.yaml", out, ROOT, quiet=True)
    assert res["exit_code"] == 0
    return out


@pytest.fixture(scope="module")
def area(tmp_path_factory):
    d = tmp_path_factory.mktemp("w10")
    return {"staging": d / "staging", "worklog": d / "worklog"}


def _start(prev_build, area, run_id, **kw):
    kw.setdefault("route", "recorded")
    kw.setdefault("cassette", CASSETTE)
    return W.start("ADD-03", PDF, pack=ROOT / "config/pack.yaml", evidence=prev_build, staging=area["staging"],
                   worklog=area["worklog"], run_id=run_id, batch_size=8, echo=quiet, sleep=lambda s: None, **kw)


@pytest.fixture(scope="module")
def full(prev_build, area):
    """One full run (background pre-addendum build, cache on), with the real state fingerprinted around it."""
    before = tree_files(*REAL)
    with WriteAudit() as audit:
        res = _start(prev_build, area, "full")
    after = tree_files(*REAL)
    cp = W.load("full", area["staging"])
    return {"res": res, "cp": cp, "dir": Path(res["run_dir"]), "before": before, "after": after, "audit": audit}


def _yaml(p: Path):
    return yaml.safe_load(p.read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------------------------- the full run

def test_a_recorded_run_ends_partial_with_every_provision_listed(full):
    res, cp, d = full["res"], full["cp"].data, full["dir"]
    assert res["status"] == "partial" and res["exit_code"] == 0, res
    assert "unresolved" in res["status_reason"]
    provs = cp["provisions"]
    assert len(provs) == 42                                         # every provision of the addendum
    assert {p for p, v in provs.items() if v["status"] == "validated"} == {
        "ADD-03:cover/para1", "ADD-03:1.2", "ADD-03:2.1", "ADD-03:2.4", "ADD-03:3.1", "ADD-03:7.1", "ADD-03:7.2"}
    assert {p for p, v in provs.items() if v["status"] == "unaccounted"} == {
        "ADD-03:cover/para2", "ADD-03:cover/para3", "ADD-03:1.1", "ADD-03:2.2", "ADD-03:2.3", "ADD-03:3.2", "ADD-03:3.3"}
    assert sum(v["status"] == "pending" for v in provs.values()) == 28   # their batches are not recorded
    for v in provs.values():                                       # pending -> proposed -> validated, in order
        assert [h["status"] for h in v["history"]] in (["pending"], ["pending", "unaccounted"],
                                                        ["pending", "proposed", "validated"])
    assert {x["kind"] for x in cp["structure"]} == {"heading", "table"}  # the non-provision units are listed too
    # session 13 (implementer D, deliberate): the analysis batches are planned by structure (batching.plan_structured):
    # Appendix A's paragraph joins the rows and notes of the table printed under its heading (T2-2-rev, one structure
    # of 13 kept whole within the token budget), so the plan has 6 batches where the count plan had 7 (Q15-Q21 with
    # AppA/para1, then T2-2-rev split 8 + 4); the three batches the cassette records are unchanged
    assert [cp["batches"][k]["status"] for k in sorted(cp["batches"]) if k.startswith("analysis")] == \
        ["done", "done", "failed", "done", "failed", "failed"]
    for s in W.STEPS:                                              # every step done, with its wall-clock time
        assert cp["steps"][s]["status"] == "done" and cp["steps"][s]["seconds"] >= 0, s
    assert sum(cp["steps"][s]["seconds"] for s in W.STEPS) < 30 * 60
    # the candidate op file accounts for every provision: 3 ops, 2 no_effect, the rest unresolved with the reason
    of = _yaml(d / "candidate/curation/amendments/ADD-03.yaml")
    assert {o["id"] for o in of["ops"]} == {"ADD-03/2.1", "ADD-03/2.4", "ADD-03/3.1"}
    assert all(o["review"] == "proposed" and o["origin"] == "assistant" and "AI workflow run full" in o["note"]
               for o in of["ops"])
    disp = {x["provision"]: x for x in of["dispositions"]}
    assert {p for p, x in disp.items() if x["disposition"] == "no_effect"} == {"ADD-03:cover/para1", "ADD-03:1.2"}
    assert len(disp) + 3 == 42
    assert "not analysed: batch analysis-003 failed" in disp["ADD-03:4.1"]["reason"]
    assert "no item proposed" in disp["ADD-03:2.2"]["reason"] and "escalated" in disp["ADD-03:7.2"]["reason"]
    md = (d / "review/index.md").read_text(encoding="utf-8")
    assert all(f"### {p} " in md for p in provs)                  # the review lists the chain of every provision
    first = md.index("## First: unresolved provisions and escalations")
    assert first < md.index("## Per provision")                   # the unresolved and escalated come first
    assert (d / "review/index.html").read_text(encoding="utf-8").count(BANNER) >= 1
    assert "## Timings" in md and "target 30 min" in md


def test_the_chain_from_evidence_to_output_difference(full):
    d = full["dir"]
    md = (d / "review/index.md").read_text(encoding="utf-8")
    sec = md[md.index("### ADD-03:3.1 "):md.index("### ADD-03:3.2 ")]
    assert "source: “In Volume I Clause 7.1" in sec                # source evidence
    assert "`ADD-03/3.1` amendment_op replace_text VOL-I:7.1" in sec  # proposed transition
    assert "validation: **evidence_verified**" in sec               # validation
    assert "units changed VOL-I:7.1" in sec and "VOL-I-7.1-01" in sec  # downstream impact
    assert "downstream proposal `D1` row_reading" in sec
    assert "output difference:" in sec and "CHANGED VOL-I-7.1-01" in sec
    diff = json.loads((d / "review/diff.json").read_text(encoding="utf-8"))
    assert diff["from"] == "ADD-02" and diff["to"] == "ADD-03" and "ADD-03-6.8-01" in diff["requirements"]["new"]


def test_candidate_outputs_exist_with_the_banner_and_the_a1_status_column(full):
    d, cp = full["dir"], full["cp"].data
    out = d / "candidate/out"
    assert cp["steps"]["outputs"]["exit_code"] == 0 and not cp["steps"]["outputs"]["refused"]
    for f in ("a1/a1.xlsx", "a1/a1.csv", "a2/a2.md", "a3/a3.pdf", "a3/a3_detail.html", "a4/clarification_register.md",
              "a5/README.md", "a5/gantt.html", "a5/programme.csv", "README.md", "CANDIDATE.md"):
        assert (out / f).is_file(), f
    for f in ("README.md", "a2/a2.md", "a4/clarification_register.md", "a5/README.md"):
        assert (out / f).read_text(encoding="utf-8").startswith(f"> **{BANNER}**"), f
    assert BANNER in (out / "a3/a3_detail.html").read_text(encoding="utf-8")
    assert BANNER in (out / "a5/gantt.html").read_text(encoding="utf-8")
    assert (out / "a1/a1.csv").read_text(encoding="utf-8").startswith(f"# {BANNER}")
    assert BANNER in pymupdf.open(out / "a3/a3.pdf")[0].get_text()
    from openpyxl import load_workbook
    ws = load_workbook(out / "a1/a1.xlsx").active
    assert ws["A1"].value.startswith("CANDIDATE") and BANNER in ws["A2"].value
    a1 = json.loads((out / "a1/a1.json").read_text(encoding="utf-8"))
    assert a1["_candidate"] == BANNER and a1["columns"][1]["key"] == "candidate_status"
    st = {r["id"]: r["candidate_status"] for r in a1["rows"]}
    assert st["ADD-03-6.8-01"].startswith("PROPOSED BY THE AI WORKFLOW: new row")
    assert st["VOL-I-7.1-01"].startswith("PROPOSED BY THE AI WORKFLOW: reading re-made")
    assert st["VOL-I-6.1-01"].startswith("UNRESOLVED: STALE")      # its re-made reading was insufficient: not promoted
    assert any(v.startswith("proposed (existing row") for v in st.values())
    assert not any(v.startswith("DECIDED") for v in st.values())    # nobody has decided anything
    # the pre-addendum outputs are the last validated state: built once, no banner, ADD-03 absent
    before = d / "candidate/out-before"
    assert (before / "a1/a1.xlsx").is_file() and not (before / "CANDIDATE.md").exists()
    assert "ADD-03" not in json.loads((before / "a1/a1.json").read_text(encoding="utf-8"))["stages"]
    assert cp["out_before"]["status"] == "done"
    cmp_ = json.loads((d / "review/outputs-before-after.json").read_text(encoding="utf-8"))
    assert "ADD-03-6.8-01" in cmp_["a1"]["new"] and "register-delivery-reps" in cmp_["a5"]["new"]


def test_promotion_goes_into_the_candidate_only(full):
    audit = full["audit"]
    assert audit.under(ROOT) == [], f"this process wrote inside the repository: {audit.under(ROOT)[:5]}"
    changed = sorted(k for k in set(full["before"]) | set(full["after"])
                     if full["before"].get(k) != full["after"].get(k))
    if changed:                                                  # only another engineer's concurrent edit can do this
        warnings.warn(f"changed by ANOTHER process during the run (this one wrote nothing there): {changed}")
    assert not (ROOT / "curation/reviews/decisions.yaml").exists()
    cp = full["cp"].data
    assert cp["inputs"]["real_hashes"]                              # the real inputs were fingerprinted at the start
    d = full["dir"] / "candidate"
    rows = _yaml(d / "curation/register/rows/ADD-03-ai.yaml")
    assert [r["id"] for r in rows["rows"]] == ["ADD-03-6.8-01"] and "PROPOSED" in rows["prepared_by"]
    reg = (d / "curation/register/rows.yaml").read_text(encoding="utf-8")
    assert "Proposal open for acceptance for 180 days from the Proposal Due Date" in reg
    assert "one hundred and eighty (180) days from the Proposal Due Date" in reg
    assert "Proposal open for acceptance for 180 days" not in (ROOT / "curation/register/rows.yaml").read_text(encoding="utf-8")
    tpl = _yaml(d / "curation/activity_templates.yaml")
    assert [a["id"] for a in tpl["EV-DELIVERY-REGISTRATION"]] == ["register-delivery-reps"]
    lt = _yaml(d / "config/assumptions.yaml")["lead_times"]["delivery_registration"]
    assert lt["basis"].startswith("PROVISIONAL ASSUMPTION")
    rel = _yaml(d / "curation/relationships.yaml")["relationships"]
    mine = [e for e in rel if str(e.get("origin", "")).startswith("ai:")]
    assert len(mine) == 1 and mine[0]["status"] == "proposed" and mine[0]["to"] == "deliver"
    clar = {c["id"]: c for c in _yaml(d / "curation/clarifications/register.yaml")["clarifications"]}
    assert clar["CQ-ADD03-DELIVERY-REPS"]["response_status"] == "draft, not sent"
    ids = _yaml(d / "curation/register/ids.yaml")["ids"]
    assert "ADD-03-6.8-01" in ids                                  # check-register --update-ids on the candidate's ledger
    assert "ADD-03-6.8-01" not in _yaml(ROOT / "curation/register/ids.yaml")["ids"]


def test_downstream_proposals_are_validated_in_the_candidate(full):
    d = full["dir"]
    data = _yaml(d / "downstream/proposals.yaml")
    st = {it["id"]: it["verification_status"] for it in data["downstream_set"]["items"]}
    # session 12 (part 1, tests/test_session12_human_owned.py): D6 is an issue, a matter kept open for people, and D7
    # asks for the response status 'sent' (tenderpack.human_owned): their quotations still verify, their conclusions are
    # a person's, so both are interpretation_pending (still promotable as proposals), no longer evidence_verified
    assert st == {"D1": "interpretation_pending", "D2": "insufficient_evidence", "D3": "interpretation_pending",
                  "D4": "evidence_verified", "D5": "interpretation_pending", "D6": "interpretation_pending",
                  "D7": "interpretation_pending", "D8": "interpretation_pending"}
    d2 = next(it for it in data["downstream_set"]["items"] if it["id"] == "D2")
    assert any(not v["ok"] and "quote not found" in v["detail"] for v in d2["validation"])   # a fabricated quotation
    ow = {(o["item"], o.get("field")) for o in data["controller"]["overwrites"]}
    assert ("D7", "entry.response_status") in ow and ("D8", "payload.status") in ow          # never sent; never confirmed
    items = full["cp"].data["downstream"]["items"]
    assert all(v["status"] == "validated" for v in items.values()) and len(items) == 8
    assert [h["status"] for h in items["downstream-001/D1"]["history"]] == ["proposed", "validated"]


def test_an_unknown_change_type_is_escalated_with_its_evidence_and_scope(full):
    d, cp = full["dir"], full["cp"].data
    assert cp["steps"]["validation"]["converted"] == [{"item": "ADD-03/7.1", "batch": "analysis-004",
                                                       "type": "reletter_items"}]
    ps = _yaml(d / "ai/full-combined/proposals.yaml")["proposal_set"]
    it = next(x for x in ps["items"] if x["id"] == "ADD-03/7.1")
    assert it["statement_type"] == "escalation" and it["verification_status"] == "escalated"
    assert "reletter_items" in it["payload"]["what_is_unsupported"] and it["evidence"]
    md = (d / "review/index.md").read_text(encoding="utf-8")
    line = md[md.index("**ESCALATED ADD-03/7.1**"):]
    line = line[:line.index("\n- **")]
    assert "never forced into a known type" in line and "ADD-03:7.1 p2" in line
    assert "affected scope: units" in line and "VOL-I-9.1-01" in line and "form-4g-sign" in line


# ---------------------------------------------------------------------------------------------- resumption

def test_a_run_killed_between_batches_resumes_without_asking_twice(prev_build, area, monkeypatch):
    calls: list[list[str]] = []
    real = W.WorkflowCassette.provider

    def killer(self, phase, keys, used):
        calls.append(list(keys))
        if phase == "analysis" and len([c for c in calls if c]) == 2:
            raise KeyboardInterrupt("simulated kill during the second batch")
        return real(self, phase, keys, used)
    monkeypatch.setattr(W.WorkflowCassette, "provider", killer)
    with pytest.raises(KeyboardInterrupt):
        _start(prev_build, area, "resumed", background_before=False, stop_after="validation")
    cp = W.load("resumed", area["staging"]).data
    b = cp["batches"]
    assert b["analysis-001"]["status"] == "done" and b["analysis-002"]["status"] == "interrupted"
    assert cp["status"] == "stopped" and "interrupted" in cp["status_reason"]
    first_run = b["analysis-001"]["staged_run"]
    staged_before = sorted(p.name for p in (Path(cp["candidate"]["dir"]).parent / "ai").iterdir() if p.is_dir())
    assert staged_before == [first_run]
    # a dead process's lock is taken over by the resume, and the takeover is recorded
    lock = Path(cp["candidate"]["dir"]).parent / "run.lock"
    lock.write_text(json.dumps({"pid": 2 ** 22 + 12345, "host": socket.gethostname(), "created": "x", "token": "t"}))
    monkeypatch.setattr(W.WorkflowCassette, "provider", real)
    calls.clear()
    res = W.resume("resumed", area["staging"], stop_after="validation", echo=quiet, sleep=lambda s: None)
    assert res["status"] == "stopped" and "stopped after validation" in res["status_reason"]
    cp = W.load("resumed", area["staging"]).data
    assert any(e["event"] == "stale_run_lock_taken_over" for e in cp["events"])
    assert cp["batches"]["analysis-001"]["staged_run"] == first_run      # not asked again
    assert cp["batches"]["analysis-002"]["status"] == "done" and cp["batches"]["analysis-002"]["attempts"] == 2
    staged = sorted(p.name for p in (Path(cp["candidate"]["dir"]).parent / "ai").iterdir() if p.is_dir())
    assert len(staged) == 4 and first_run in staged and "resumed-combined" in staged   # 001, 002, 004 and the combined
    ps = _yaml(Path(cp["candidate"]["dir"]).parent / "ai/resumed-combined/proposals.yaml")["proposal_set"]
    ids = [it["id"] for it in ps["items"]]
    assert len(ids) == len(set(ids)) == 7                         # no duplicated item
    assert sorted(ids) == sorted(["ADD-03/cover/para1", "ADD-03/1.2", "ADD-03/2.1", "ADD-03/2.4", "ADD-03/3.1",
                                  "ADD-03/7.1", "ADD-03/7.2"])
    for p in ("ADD-03:cover/para1", "ADD-03:2.1"):
        assert [h["status"] for h in cp["provisions"][p]["history"]] == ["pending", "proposed", "validated"]


# ---------------------------------------------------------------------------------------------- a refused build

def test_a_refused_outputs_build_keeps_out_before(prev_build, area, full):
    """The candidate is made unusable after check-register (its evidence build edited after ingest: E01): the
    outputs step refuses before building, out-before stays the last validated state and the review says so."""
    res = _start(prev_build, area, "refused", background_before=False, stop_after="check_register")
    assert res["status"] == "stopped"
    d = Path(res["run_dir"])
    units = d / "candidate/build/units.json"
    units.write_text(units.read_text(encoding="utf-8") + " ", encoding="utf-8")
    res = W.resume("refused", area["staging"], retry_failed=False, echo=quiet, sleep=lambda s: None)
    cp = W.load("refused", area["staging"]).data
    o = cp["steps"]["outputs"]
    assert o["refused"] and o["exit_code"] == 2 and o["reason"].startswith("pre-flight: E01")
    assert res["status"] == "partial" and res["exit_code"] == 1 and "refused" in res["status_reason"]
    assert not (d / "candidate/out").exists()
    assert (d / "candidate/out-before/a1/a1.xlsx").is_file() and cp["out_before"]["status"] == "done"
    assert cp["out_before"]["result"].get("from_cache")            # built once by the full run, reused here
    md = (d / "review/index.md").read_text(encoding="utf-8")
    assert "**The candidate outputs build was REFUSED** (exit 2)" in md
    assert "The last validated state stays `../candidate/out-before/`." in md


# ---------------------------------------------------------------------------------------------- the manual host path

def test_the_manual_host_path_waits_records_the_submission_and_continues(prev_build, area, tmp_path):
    res = _start(prev_build, area, "host", route="host", cassette=None, host_mode="manual", background_before=False)
    assert res["status"] == "waiting_for_host" and res["exit_code"] == 4 and res["waiting"] == ["analysis-001"]
    cp = W.load("host", area["staging"]).data
    packet = json.loads(Path(cp["batches"]["analysis-001"]["packet"]).read_text(encoding="utf-8"))
    assert packet["workflow"]["answer_only"] == cp["batches"]["analysis-001"]["provisions"]
    assert "submit-batch host" in packet["workflow"]["submit_with"]
    turn = _yaml(CASSETTE)["sessions"][0]["turns"][0]["response"]["text"]
    sub = tmp_path / "batch1.json"
    sub.write_text(turn.replace("${state}", json.dumps(packet["state"])), encoding="utf-8")
    with pytest.raises(Exception, match="--by"):
        W.submit_batch("host", sub, " ", "host-declared-model", staging=area["staging"], echo=quiet)
    res = W.submit_batch("host", sub, "Fixture Host Session", "host-declared-model", staging=area["staging"], echo=quiet)
    assert res["status"] == "waiting_for_host" and res["waiting"] == ["analysis-002"]
    cp = W.load("host", area["staging"]).data
    iv = cp["interventions"]
    assert len(iv) == 1 and iv[0]["kind"] == "submit-batch" and iv[0]["by"] == "Fixture Host Session"
    assert iv[0]["batch"] == "analysis-001" and len(iv[0]["sha256"]) == 64 and iv[0]["items"] == 2
    assert cp["provisions"]["ADD-03:1.2"]["status"] == "validated"
    assert cp["provisions"]["ADD-03:1.1"]["status"] == "unaccounted"
    staged = _yaml(Path(cp["candidate"]["dir"]).parent / "ai" / cp["batches"]["analysis-001"]["staged_run"] /
                   "proposals.yaml")["proposal_set"]
    assert staged["route"] == "host" and staged["model_requested"] == "host-declared-model"


# ---------------------------------------------------------------------------------------------- image regions (readings)

def _addendum_with_image(d: Path) -> Path:
    """Blind-02's Addendum No. 3 with an image of two printed English lines added on page 3 (no text layer): an
    addendum ingest refuses (C05 unread) until the region has a reading."""
    doc = pymupdf.open(PDF)
    src = pymupdf.open()
    pg = src.new_page(width=460, height=80)
    pg.insert_text((12, 30), "Note: Bidders shall also submit one additional USB copy of Envelope B.", fontname="helv",
                   fontsize=11)
    pg.insert_text((12, 58), "The additional USB copy shall be encrypted as Clause 6.5 requires.", fontname="helv",
                   fontsize=11)
    doc[2].insert_image(pymupdf.Rect(66, 240, 526, 320), pixmap=pg.get_pixmap(dpi=200))
    out = Path(d) / "ADD-03_with_image.pdf"
    doc.save(out)
    return out


@pytest.fixture(scope="module")
def image_pdf(tmp_path_factory):
    return _addendum_with_image(tmp_path_factory.mktemp("w10-img"))


GOOD_LINES = ["Note: Bidders shall also submit one additional USB copy of Envelope B.",
              "The additional USB copy shall be encrypted as Clause 6.5 requires."]


def _reading(source: dict, lines=GOOD_LINES) -> dict:
    return {"region_id": "ADD-03-p3-r1", "unit_id": "ADD-03:p3-image", "title": "Note on page 3 (image)",
            "source": source, "content_type": "text", "languages": ["en"], "prepared_by": "test", "method": "test",
            "blocks": [{"key": "note", "lang": "en", "role": "note",
                        "lines": [{"band": i, "source": s} for i, s in enumerate(lines)]}]}


def test_an_image_region_gets_a_proposed_reading_pending_review_and_ingest_runs_again(prev_build, area, image_pdf):
    res = W.start("ADD-03", image_pdf, pack=ROOT / "config/pack.yaml", evidence=prev_build, staging=area["staging"],
                  worklog=area["worklog"], run_id="reading", route="recorded", cassette=CASSETTES / "workflow_reading.yaml",
                  background_before=False, stop_after="readings", echo=quiet, sleep=lambda s: None)
    assert res["status"] == "stopped" and "stopped after readings" in res["status_reason"], res
    cp = W.load("reading", area["staging"]).data
    ing, rd = cp["steps"]["ingest"], cp["steps"]["readings"]
    assert ing["exit_code"] == 2 and ing["refused_for_readings"] and ing["unread_regions"] == ["ADD-03-p3-r1"]
    assert rd["status"] == "done" and rd["reingest"]["exit_code"] == 0
    b = cp["batches"]["reading-ADD-03-p3-r1"]
    assert b["status"] == "done" and b["verification_status"] == "interpretation_pending"   # never higher
    d = Path(res["run_dir"])
    f = d / "candidate/curation/readings/ADD-03-p3-r1.yaml"
    assert Path(b["reading_file"]) == f and not (ROOT / "curation/readings/ADD-03-p3-r1.yaml").exists()
    text = f.read_text(encoding="utf-8")
    assert text.startswith("# AI-PROPOSED READING — PENDING HUMAN REVIEW. Not approved.")
    r = yaml.safe_load(text)
    assert r["prepared_by"].startswith("AI-assisted: recorded route") and "AI workflow run reading" in r["prepared_by"]
    assert r["blocks"][0]["lines"][1]["source"] == GOOD_LINES[1] and r["source"]["page"] == 3
    approvals = (d / "candidate/curation/approvals.yaml").read_text(encoding="utf-8")
    assert "ADD-03-p3-r1" not in approvals and approvals == (ROOT / "curation/approvals.yaml").read_text(encoding="utf-8")
    # the reading's units are provisions of the addendum now, pending a person
    assert {"ADD-03:p3-image/note"} <= set(cp["provisions"]) and cp["provisions"]["ADD-03:p3-image/note"]["status"] == "pending"
    units = json.loads((d / "candidate/build/units.json").read_text(encoding="utf-8"))["units"]
    u = next(x for x in units if x["unit_id"] == "ADD-03:p3-image/note")
    assert u["origin"] == "image_reading" and u["reading"]["status"] == "pending"
    assert "PENDING" in (d / "candidate/build/review/ADD-03-p3-r1/packet.md").read_text(encoding="utf-8")
    assert not (d / "candidate/build.failed").exists()
    res2 = json.loads(Path(b["result"]).read_text(encoding="utf-8"))
    assert {"field": "reading.prepared_by", "proposer_value": "the proposer (overwritten by the controller)"} in \
        res2["controller"]["overwrites"]
    log = [json.loads(x) for x in (d / "ai/reading-reading-ADD-03-p3-r1/log.jsonl").read_text().splitlines()]
    call = next(e for e in log if e["event"] == "tool_call")
    assert call["name"] == "get_region" and call["ok"] and len(call["images"]) == 3   # native, context, crop sent
    md = W.review_markdown(W.Ctx(W.load("reading", area["staging"]), quiet))
    sec = md[md.index("## Readings of the addendum's image regions"):]
    assert "**ADD-03-p3-r1** — done; controller status **interpretation_pending**" in sec
    assert "candidate/build/review/ADD-03-p3-r1/packet.html" in sec and "PENDING HUMAN REVIEW" in md


def test_an_unusable_reading_stops_the_run_with_the_reason(prev_build, area, image_pdf):
    res = W.start("ADD-03", image_pdf, pack=ROOT / "config/pack.yaml", evidence=prev_build, staging=area["staging"],
                  worklog=area["worklog"], run_id="reading-bad", route="recorded",
                  cassette=CASSETTES / "workflow_reading_bad.yaml", background_before=False, echo=quiet,
                  sleep=lambda s: None)
    assert res["status"] == "stopped" and "no usable reading for the image region(s) ['ADD-03-p3-r1']" in res["status_reason"]
    assert "RD4" in res["status_reason"] and "insufficient_evidence" in res["status_reason"]
    cp = W.load("reading-bad", area["staging"]).data
    assert cp["batches"]["reading-ADD-03-p3-r1"]["status"] == "failed"
    d = Path(res["run_dir"])
    assert not (d / "candidate/curation/readings/ADD-03-p3-r1.yaml").exists()
    assert (d / "candidate/build.failed/regions/ADD-03-p3-r1/native.png").is_file()      # kept for the next attempt


def test_the_region_tools_read_a_refused_build(prev_build, area, image_pdf):
    """get_region and validate_reading over the refused build of the run above (and over MCP, with the images)."""
    from tenderpack.ai.tools import ToolError, Workspace, call_tool
    from tenderpack.mcp_server import Server
    d = Path(W.load("reading-bad", area["staging"]).data["candidate"]["dir"])
    ws = Workspace(d / "build.failed", d / "pack.yaml", ROOT, area["staging"], area["worklog"])
    r = call_tool(ws, "get_region", {"region_id": "ADD-03-p3-r1"}, caller="mcp")
    assert (r["doc"], r["page"], r["kind"]) == ("ADD-03", 3, "image") and r["grid"] is None
    assert [b["index"] for b in r["bands"] if b["kind"] == "text"] == [0, 1] and r["text_bands"] == 2
    assert [c["kind"] for c in r["crops"]] == ["native", "context", "region_crop"]
    assert r["crops"][0]["sha256"] == r["native"]["sha256"]
    rb = call_tool(ws, "get_region", {"region_id": "ADD-03-p3-r1", "bands": [1]}, caller="mcp")
    assert [(c["kind"], c["band"], c["side"]) for c in rb["crops"]] == [("band", 1, "full")]
    with pytest.raises(ToolError, match="no region"):
        call_tool(ws, "get_region", {"region_id": "ADD-03-p9-r9"}, caller="mcp")
    with pytest.raises(ToolError, match="no tool"):
        call_tool(ws, "get_region", {"region_id": "ADD-03-p3-r1"}, caller="model")    # not offered to the proposer
    src = {"doc": "ADD-03", "page": 3, "bbox_pt": r["bbox_pt"], "native_sha256": r["native"]["sha256"]}
    ok = call_tool(ws, "validate_reading", {"reading": _reading(src)}, caller="mcp")
    assert ok["parsed"] and ok["ok"] and not ok["errors"]
    short = call_tool(ws, "validate_reading", {"reading": _reading(src, GOOD_LINES[:1])}, caller="mcp")
    assert not short["ok"] and any(e.startswith("RD4") and "(1, 'full')" in e for e in short["errors"])
    other = call_tool(ws, "validate_reading", {"reading": _reading(dict(src, native_sha256="0" * 64))}, caller="mcp")
    assert any(e.startswith("RD2") for e in other["errors"])
    bad = call_tool(ws, "validate_reading", {"reading": {"region_id": "x"}}, caller="mcp")
    assert not bad["parsed"] and "not a Reading" in bad["errors"][0]
    # over MCP the images are sent with the result
    out = Server(ws).handle({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                             "params": {"name": "get_region", "arguments": {"region_id": "ADD-03-p3-r1"}}})
    content = out["result"]["content"]
    assert not out["result"]["isError"] and [c["type"] for c in content] == ["text", "image", "image", "image"]
    assert json.loads(content[0]["text"])["images_attached"][0]["kind"] == "native"
    # a crop that is not the file the build recorded is never served
    crop = d / "build.failed/regions/ADD-03-p3-r1/context.png"
    saved = crop.read_bytes()
    try:
        crop.write_bytes(saved + b"x")
        with pytest.raises(ToolError, match="integrity failure"):
            call_tool(ws, "get_region", {"region_id": "ADD-03-p3-r1"}, caller="mcp")
    finally:
        crop.write_bytes(saved)
