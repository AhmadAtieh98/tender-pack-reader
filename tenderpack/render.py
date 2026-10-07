"""Output writers for the deliverables: A1 (JSON, CSV, Excel), plain CSV/JSON tables (A2, A5) and
the one-page A3 PDF. Presentation only: every writer takes plain data and adds nothing to it.

Determinism. Writing the same input twice, in any directory, at any time, gives identical bytes.
  JSON  util.dump_json (sorted keys, fixed separators).
  CSV   UTF-8 with BOM (so Excel reads Arabic correctly), "\\r\\n" line ends, every value as text.
  XLSX  openpyxl stamps the save time into docProps/core.xml and into every zip entry. The
        workbook properties are fixed, core.xml is serialised again after saving, and the zip is
        rewritten with fixed entry times, attributes and order, using the same (deflate) compression.
  PDF   fixed metadata and no /ID; fonts come from PyMuPDF's bundled set (never system fonts) and
        are subset by MuPDF, whose subset prefixes depend only on the content.

Arabic in A3. MuPDF 1.28 ignores dir="rtl" on an inline <span> (checked by rendering), so the span
also carries an RLE...PDF embedding. Latin expressions inside Arabic text ("4-C") are embedded left to
right (arabic.display_form), and LRM marks keep the neighbouring quotes and separators in the
left-to-right line. Text whose first strong character is Arabic is laid out right to left as a
whole; English text keeps its order and only its Arabic runs are laid out right to left. The
controls exist only in the HTML handed to MuPDF; stored text is never changed.
"""
from __future__ import annotations

import csv
import html
import io
import math
import re
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.xml.functions import tostring

from .arabic import display_form
from .util import dump_json, r1

FIXED_TIME = datetime(2026, 10, 1, 0, 0, 0)       # workbook created/modified
PDF_DATE = "D:20261001000000Z"
PRODUCER = "tenderpack"
ZIP_TIME = (1980, 1, 1, 0, 0, 0)                   # earliest time a zip entry can carry
A1_SHEET = "A1 register"
A1_HEADER_ROW = 4                                  # rows 1-2: banner; row 3: blank
A3_SCALE_LOW = 0.9                                 # 8.5 pt body text never drops below 7.6 pt; else: prioritise
A3_MIN_TEXT_PT = 7.5                               # smallest rendered text allowed on A3 (C43)
BANNER_SCREEN_WIDTH = 160                          # column-width units a sheet's banner is merged across (about a screen)
BANNER_LINE_PT = 15                                # row height per wrapped line of the default 11 pt font


# ------------------------------------------------------------------ values as text

def cell_text(value: Any) -> str:
    """A value as CSV text: None -> "", bool -> TRUE/FALSE, lists joined with "; "."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (list, tuple)):
        return "; ".join(cell_text(v) for v in value)
    if isinstance(value, (str, int, float)):
        return str(value)
    # anything else (dict, set, ...) has no agreed text form, and a set's order is not stable
    raise TypeError(f"unsupported cell value of type {type(value).__name__}: {value!r}")


def _xlsx_value(value: Any) -> Any:
    """Numbers and booleans stay typed in Excel; lists are joined like in the CSV."""
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return cell_text(value)


def _columns(table: dict) -> list[dict]:
    cols = table["columns"]
    keys = [c["key"] for c in cols]
    dup = sorted({k for k in keys if keys.count(k) > 1})
    if dup:
        raise ValueError(f"duplicate column keys: {dup}")
    return cols


# ------------------------------------------------------------------ CSV and JSON

def _csv_bytes(table: dict) -> bytes:
    cols = _columns(table)
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([c["header"] for c in cols])
    for row in table["rows"]:
        w.writerow([cell_text(row.get(c["key"])) for c in cols])
    return buf.getvalue().encode("utf-8-sig")


def write_csv_json(table: dict, out_dir: Path, stem: str) -> list[Path]:
    """`stem.csv` (declared columns only) and `stem.json` (the whole table) in out_dir."""
    out_dir = Path(out_dir)
    data = _csv_bytes(table)                       # fails before anything is written
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = [out_dir / f"{stem}.csv", out_dir / f"{stem}.json"]
    paths[0].write_bytes(data)
    dump_json(table, paths[1])
    return paths


# ------------------------------------------------------------------ A1: Excel

STATUS_FILLS = (                                   # first match wins
    (lambda v: v.startswith(("DELETED", "REMOVED")), "D9D9D9"),
    (lambda v: v.startswith(("REPLACED", "REVOKED")), "DDEBF7"),      # session 11 audit (A1-11): out of force, replaced
    (lambda v: "STALE" in v, "FCE4D6"),
    (lambda v: v.startswith(("NEW", "REINSTATED")), "E2EFDA"),
    (lambda v: v.startswith("AMENDED"), "FFF2CC"),
    (lambda v: v.startswith(("NOT ISSUED", "NOT IN FORCE")), "F2F2F2"),
)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
HEADER_BORDER = Border(bottom=Side(style="thin"))


def status_fill(value: Any) -> str | None:
    """Fill colour (RRGGBB) of a per-stage status cell, or None."""
    text = cell_text(value).strip()
    return next((rgb for test, rgb in STATUS_FILLS if test(text)), None)


def _is_status_column(col: dict) -> bool:
    # per-stage status columns: header "status:..." as specified; key "status:..." when the header is prose
    return col["header"].startswith("status:") or col["key"].startswith("status:")


def _write_table(ws, cols: list[dict], rows: list[dict], header_row: int, styled: bool) -> None:
    """Header, rows, widths, frozen header and autofilter. `styled`: status fills and pending italics."""
    for c, col in enumerate(cols, 1):
        cell = ws.cell(header_row, c, col["header"])
        cell.font, cell.alignment, cell.border = Font(bold=True), WRAP_TOP, HEADER_BORDER
        ws.column_dimensions[get_column_letter(c)].width = col["width"]
    for r, row in enumerate(rows, header_row + 1):
        for c, col in enumerate(cols, 1):
            value = row.get(col["key"])
            cell = ws.cell(r, c, _xlsx_value(value))
            if isinstance(cell.value, str):
                cell.data_type = "s"               # text such as "=..." or "#N/A" stays text
            cell.alignment = WRAP_TOP
            if not styled:
                continue
            if _is_status_column(col) and (rgb := status_fill(value)):
                cell.fill = PatternFill(fill_type="solid", fgColor=rgb)
            if "status" in col["key"].lower() and "pending" in cell_text(value).lower():
                cell.font = Font(italic=True)
    last = get_column_letter(max(len(cols), 1))
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
    ws.auto_filter.ref = f"A{header_row}:{last}{header_row + len(rows)}"


def _banner(ws, row: int, text: str, cols: list[dict]) -> None:
    """A long notice in one cell, readable in the sheet (session 12, audit R-5: the A1 notice, over 1,100 characters,
    sat in one unwrapped, unmerged cell at the default row height and could be read only in the formula bar). The cell
    is merged across the leading columns that fit about one screen (at least two, never wider than the table), wrapped
    at the top, and its row given the height of the wrapped lines plus one (Excel does not size a merged row itself;
    a column-width unit is about one character, so the line count errs on the generous side). Only that row changes."""
    widths = [c["width"] for c in cols] or [BANNER_SCREEN_WIDTH]
    n = total = 0
    for w in widths:
        if n >= 2 and total + w > BANNER_SCREEN_WIDTH:
            break
        n, total = n + 1, total + w
    cell = ws.cell(row, 1, text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if n > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
    lines = sum(max(1, math.ceil(len(part) / max(total, 1))) for part in str(text).split("\n"))
    ws.row_dimensions[row].height = min(409, BANNER_LINE_PT * (lines + 1))


def _a1_workbook(a1: dict) -> Workbook:
    wb = Workbook()
    p = wb.properties
    p.creator = p.lastModifiedBy = PRODUCER
    p.created = p.modified = FIXED_TIME
    p.title = a1["title"]
    ws = wb.active
    ws.title = A1_SHEET
    ws["A1"] = a1["title"]
    ws["A1"].font = Font(bold=True, size=12)
    _banner(ws, 2, a1["notice"], _columns(a1))
    ws["A2"].font = Font(italic=True)
    _write_table(ws, _columns(a1), a1["rows"], A1_HEADER_ROW, styled=True)
    for name, sheet in a1.get("sheets", {}).items():
        if name == A1_SHEET:
            raise ValueError(f"extra sheet may not be named {A1_SHEET!r}")
        sh = wb.create_sheet(name)
        if sheet.get("notice"):                    # session 12 (R-6): a legend above the table, header one row down
            _banner(sh, 1, sheet["notice"], _columns(sheet))
        _write_table(sh, _columns(sheet), sheet["rows"], 2 if sheet.get("notice") else 1, styled=False)
    return wb


def _xlsx_bytes(wb: Workbook) -> bytes:
    """Save, then rewrite the zip so that nothing depends on when or where it was written."""
    raw = io.BytesIO()
    wb.save(raw)                                   # stamps properties.modified with the time now
    wb.properties.modified = FIXED_TIME
    core = tostring(wb.properties.to_tree())       # serialised exactly as openpyxl's writer does
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        names = sorted(src.namelist(), key=lambda n: (n != "[Content_Types].xml", n))
        for name in names:
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0                 # not the writing platform's
            info.external_attr = 0
            dst.writestr(info, core if name == "docProps/core.xml" else src.read(name))
    return out.getvalue()


def write_a1(a1: dict, out_dir: Path) -> list[Path]:
    """`a1.json`, `a1.csv` (main rows) and `a1.xlsx` (register plus extra sheets) in out_dir."""
    out_dir = Path(out_dir)
    csv_data = _csv_bytes(a1)                      # build everything before writing anything
    xlsx_data = _xlsx_bytes(_a1_workbook(a1))
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = [out_dir / "a1.json", out_dir / "a1.csv", out_dir / "a1.xlsx"]
    dump_json(a1, paths[0])
    paths[1].write_bytes(csv_data)
    paths[2].write_bytes(xlsx_data)
    return paths


# ------------------------------------------------------------------ A3: one-page PDF

class A3OverflowError(ValueError):
    """A3 does not fit on one A4 page at scale >= A3_SCALE_LOW. Nothing is written: the content
    has to be prioritised by a person, never dropped or moved to a second page."""


LRM, RLE, POP = "\u200e", "\u202b", "\u202c"
SEP = " \u00b7 "                                  # between source, confidence and flags
_ARABIC = re.compile("[\u0600-\u06ff]")
_ARABIC_RUN = re.compile("[\u0600-\u06ff](?:[^A-Za-z]*[\u0600-\u06ff])?")   # no Latin letter inside

A3_PAGE = pymupdf.paper_rect("a4")
A3_MARGIN_X, A3_MARGIN_Y, A3_FOOTER_H = 26, 17, 26
A3_CSS = """
body {font-family: sans-serif; font-size: 8.5px; line-height: 1.1; color: #000}
.title {font-size: 13px; font-weight: bold; margin: 0 0 1px 0}
.sub {color: #333; margin: 0 0 3px 0}
.banner {background-color: #e3e3e3; border: 0.5px solid #8c8c8c; padding: 2px 5px; font-weight: bold;
         margin: 0 0 4px 0}
h2 {font-size: 9.5px; font-weight: bold; border-bottom: 0.5px solid #8c8c8c; margin: 3.5px 0 1.5px 0;
    padding: 0 0 1px 0}
p {margin: 0 0 1.5px 0}
.note {font-style: italic; color: #444}
.it {padding-left: 9px; text-indent: -9px}
.meta {color: #444}
.flag {color: #b00000; font-size: 8.5px}
.ids {font-size: 8.5px; color: #222}
a {color: #000; text-decoration: none}
.none {color: #444; font-style: italic}
.nw {white-space: nowrap}
.grp {margin: 0.5px 0 0 0}
.it2 {padding-left: 18px; text-indent: -9px; margin: 0 0 0.5px 0}
"""
A3_FOOTER_CSS = """
body {font-family: sans-serif; font-size: 8.5px; line-height: 1.2; color: #444}
.foot {border-top: 0.5px solid #8c8c8c; padding-top: 2px}
"""


def _esc(text: str) -> str:
    return html.escape(text, quote=False)


def _first_strong_rtl(text: str) -> bool:
    for ch in text:
        b = unicodedata.bidirectional(ch)
        if b in ("R", "AL"):
            return True
        if b == "L":
            return False
    return False


def _rtl(text: str) -> str:
    """One right-to-left run inside the left-to-right page (see module docstring)."""
    return f'{LRM}<span dir="rtl">{RLE}{_esc(display_form(text))}{POP}</span>{LRM}'


_ID_TOKEN = re.compile(r"\b[A-Z][A-Z0-9]*(?:-[A-Za-z0-9.]+)+(?:/[A-Za-z0-9.\-]+)*(?:\(\d+\))?")


def _nw(escaped: str) -> str:
    """Ids inside running text (VOL-I-6.1-01, I-A5-FEASIBILITY, ADD-02/T1-1-rev/note(2)) kept on one line (session 11,
    audit R-6): a line may break between ids, never at an id's internal hyphen. The span has no effect where the CSS has
    no `nw` class."""
    return _ID_TOKEN.sub(lambda m: f'<span class="nw">{m.group(0)}</span>', escaped)


def _rich(text: str) -> str:
    """Escaped HTML for one field, with any Arabic laid out right to left and ids unbroken."""
    text = text or ""
    if not _ARABIC.search(text):
        return _nw(_esc(text))
    if _first_strong_rtl(text):
        return _rtl(text)
    out, pos = [], 0
    for m in _ARABIC_RUN.finditer(text):
        out += [_nw(_esc(text[pos:m.start()])), _rtl(m.group(0))]
        pos = m.end()
    return "".join(out) + _nw(_esc(text[pos:]))


def _id(i: str, bold: bool = True) -> str:
    """An id linked to its line on a3_detail.html, never broken across lines at its internal hyphens (session 11, audit
    R-6): a list of ids breaks only after its ', ' or '; ' separators."""
    a = f'<a href="a3_detail.html#{_esc(i)}">{_esc(i)}</a>'
    return f'<span class="nw">{"<b>" + a + "</b>" if bold else a}</span>'


def _a3_item(it: dict) -> str:
    """Bold id (linked to a3_detail.html), text, computed date facts, class: "consequence" (its translation, labelled
    from the approval state), the rows that corroborate it, source, confidence, flags in red. A compact item leaves the
    quotation to the detail page."""
    parts = [f'{_id(it["id"])} {_rich(it["text"])}']
    if it.get("facts"):                                  # session 11 (A3-7, A3-9): computed dates, never typed
        parts.append("\u2014 " + "; ".join(_rich(f) for f in it["facts"]))
    if it.get("consequence") and not it.get("compact"):
        cls = f"<i>{_esc(it['class'])}</i>: " if it.get("class") else ""
        parts.append(f"\u2014 {cls}\u201c{_rich(it['consequence'])}\u201d")
        if it.get("gloss"):
            parts.append(f"({_esc(it.get('gloss_label') or 'proposed translation, not reviewed')}: \u2018{_esc(it['gloss'])}\u2019)")
    meta = [_rich(it["source"])] if it.get("source") else []
    if it.get("confidence"):
        meta.append(f"confidence {_rich(it['confidence'])}")
    if it.get("owner"):
        meta.append(f"owner: {_rich(it['owner'])}")
    if meta:
        parts.append(f'<span class="meta">[{SEP.join(meta)}]</span>')
    if it.get("flags"):
        parts.append(f'<span class="flag">{SEP.join(_rich(f) for f in it["flags"])}</span>')
    for c in it.get("corroborated_by") or []:            # session 11 (A3-6): one trigger, its corroborating sources
        parts.append("\u2014 also stated in " + _id(c["id"], bold=False) + f" ({_rich(c['source'])}"
                     + (f", confidence {_esc(c['confidence'])}" if c.get("confidence") else "") + ")"
                     + "".join(f", <span class=\"nw\">{_esc(a)}</span>" for a in c.get("also") or [])
                     + (f"; {_rich(c['adds'])}" if c.get("adds") else ""))
    return f'<p class="it">{" ".join(parts)}</p>'


def _a3_html(a3: dict) -> str:
    out = [f'<div class="title">{_rich(a3["title"])}</div>']
    if a3.get("subtitle"):
        out.append(f'<div class="sub">{_rich(a3["subtitle"])}</div>')
    if a3.get("banner"):
        out.append(f'<div class="banner">{_rich(a3["banner"])}</div>')
    tail = [] if a3.get("groups") else [*([a3["missing"]] if a3.get("missing") else []), a3["unresolved"]]
    for sec in [*a3["sections"], *tail]:
        out.append(f"<h2>{_rich(sec['heading'])}</h2>")
        if sec.get("note"):
            out.append(f'<p class="note">{_rich(sec["note"])}</p>')
        if sec.get("ids"):
            own = sec.get("owners") or {}
            out.append('<p class="ids">' + ", ".join(_id(i, bold=False)
                                                     + (f' <span class="meta">({_esc(own[i])})</span>' if own.get(i) else "")
                                                     for i in sec["ids"]) + "</p>")
            continue
        if "ids" in sec and not sec.get("items"):
            continue
        # a row that corroborates another is shown on that row's line (session 11, A3-6)
        out += [_a3_item(it) for it in sec["items"] if not it.get("corroborates")] or ['<p class="none">None.</p>']
    if a3.get("groups"):
        g = a3["groups"]
        out.append(f"<h2>{_rich(g['heading'])}</h2>")
        if g.get("note"):
            out.append(f'<p class="note">{_rich(g["note"])}</p>')
        link = _id
        for grp in g["groups"]:
            qs = (f' <span class="meta">questions drafted: {", ".join(link(q) for q in grp["questions"])}</span>'
                  if grp.get("questions") else
                  f' <span class="meta">questions drafted: {grp["n_questions"]}</span>' if grp.get("n_questions") else "")
            if grp.get("unlisted"):                # session 13 (F2; audit R2-7): questions on no issue listed here
                qs += f' <span class="meta">({_rich(grp["unlisted"])})</span>'
            if "count" in grp:                     # condensed to a count (stage2.condense_a3 level 4): ids on a3_detail.html
                out.append(f'<p class="it"><b>{_rich(grp["title"])}</b> ({grp["count"]})' + qs + "</p>")
                continue
            mark = lambda i: "\u2020\u00a0" if i.get("decide") else ""  # noqa: E731
            own = lambda i: f' <span class="meta">({_rich(i["owner"])})</span>' if i.get("owner") else ""  # noqa: E731
            # session 13 (F4; audit R2-11): a folded pending issue is named with the mark: '+1: ⚑ I-X (Owner)'
            fold = lambda i: (f' <span class="meta">+{len(i["folds"])}'  # noqa: E731
                              + (": \u2691 " + _rich("; ".join(i["folds_pending"])) if i.get("folds_pending") else "")
                              + "</span>") if i.get("folds") else ""
            if any(i.get("short") for i in grp["items"]):     # session 11 (A3-1, R-3): one line each, reason and owner
                out.append(f'<p class="grp"><b>{_rich(grp["title"])}</b> ({len(grp["items"])})' + qs + "</p>")
                out += [f'<p class="it2">{mark(i)}{link(i["id"])} {_rich(i["short"])}{own(i)}{fold(i)}</p>' for i in grp["items"]]
                continue
            items = [mark(i) + link(i["id"]) + own(i) + fold(i) for i in grp["items"]]   # condensed: ids and owners
            out.append(f'<p class="it"><b>{_rich(grp["title"])}</b> ({len(grp["items"])}): ' + "; ".join(items) + qs + "</p>")
    return "\n".join(out)


def write_a3_pdf(a3: dict, path: Path) -> dict:
    """Render A3 on exactly one A4 portrait page, scaled down to at most A3_SCALE_LOW if needed.

    Returns {"pages": 1, "scale": s, "spare_pt": h}: s = 1.0 means no scaling; h is the unused height
    below the content. Raises A3OverflowError, and writes nothing, if it does not fit.
    """
    w, h = A3_PAGE.width, A3_PAGE.height
    doc = pymupdf.open()
    page = doc.new_page(width=w, height=h)
    body = pymupdf.Rect(A3_MARGIN_X, A3_MARGIN_Y, w - A3_MARGIN_X, h - A3_MARGIN_Y - A3_FOOTER_H - 4)
    # insert_htmlbox places everything or nothing: (-1, scale) when even scale_low does not fit
    try:
        spare, scale = page.insert_htmlbox(body, _a3_html(a3), css=A3_CSS, scale_low=A3_SCALE_LOW)
    except AssertionError:          # session 12 (F5): PyMuPDF asserts 'scale >= scale_low' when the content overflows
        spare, scale = -1, 0.0      # by a rounding hair (scale 0.8999...) instead of returning -1: the same overflow
    if spare < 0 or scale < A3_SCALE_LOW or round(8.5 * scale, 2) < A3_MIN_TEXT_PT:
        raise A3OverflowError(f"A3 content does not fit on one A4 page at scale >= {A3_SCALE_LOW} "
                              f"({sum(len(s['items']) for s in a3['sections'])} items, "
                              f"{len((a3.get('unresolved') or {}).get('items') or [])} unresolved); shorten or prioritise it")
    foot = pymupdf.Rect(A3_MARGIN_X, h - A3_MARGIN_Y - A3_FOOTER_H, w - A3_MARGIN_X, h - A3_MARGIN_Y)
    fspare, _ = page.insert_htmlbox(foot, f'<div class="foot">{_rich(a3.get("footer", ""))}</div>',
                                    css=A3_FOOTER_CSS, scale_low=A3_SCALE_LOW)
    if fspare < 0:
        raise A3OverflowError(f"A3 footer does not fit in {A3_FOOTER_H} pt at scale >= {A3_SCALE_LOW}")
    assert doc.page_count == 1
    doc.subset_fonts()
    doc.set_metadata({"title": a3["title"], "creator": PRODUCER, "producer": PRODUCER,
                      "creationDate": PDF_DATE, "modDate": PDF_DATE})
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = doc.tobytes(garbage=3, deflate=True, no_new_id=True)
    doc.close()
    # Relative hrefs come out as /Launch actions, which many viewers block. The links only become visible after
    # a save, so reopen and rewrite them as /URI actions (resolved relative to the PDF), then write once.
    doc = pymupdf.open("pdf", data)
    page = doc[0]

    def launch_links():
        return [ln for ln in page.get_links() if ln.get("kind") in (pymupdf.LINK_LAUNCH, pymupdf.LINK_GOTOR)
                and "/Launch" in doc.xref_object(ln["xref"])]
    todo = [(ln["from"], (ln.get("file") or "").replace("%23", "#")) for ln in launch_links()]
    while launch_links():                                  # one at a time: deleting invalidates the list
        page.delete_link(launch_links()[0])
    for rect, target in todo:
        page.insert_link({"kind": pymupdf.LINK_URI, "from": rect, "uri": target})
    path.write_bytes(doc.tobytes(garbage=3, deflate=True, no_new_id=True))
    doc.close()
    return {"pages": 1, "scale": round(scale, 3), "spare_pt": r1(spare), "min_text_pt": round(8.5 * scale, 2)}


# ------------------------------------------------------------------ candidate A3 (session 11): a PARTIAL addendum

CANDIDATE_CSS = """
body {font-family: sans-serif; font-size: 9px; line-height: 1.2; color: #000}
.title {font-size: 14px; font-weight: bold; margin: 0 0 2px 0}
.sub {color: #333; margin: 0 0 3px 0}
.banner {background-color: #fde2e2; border: 0.8px solid #b00000; color: #7a0000; padding: 3px 6px; font-weight: bold;
         margin: 0 0 6px 0}
h2 {font-size: 11px; font-weight: bold; border-bottom: 0.5px solid #8c8c8c; margin: 8px 0 3px 0; padding: 0 0 1px 0}
h3 {font-size: 9.5px; font-weight: bold; margin: 5px 0 2px 0}
p {margin: 0 0 2px 0}
.it {padding-left: 10px; text-indent: -10px}
.sub2 {padding-left: 20px; text-indent: -10px; color: #333}
.meta {color: #444}
.flag {color: #b00000}
.mark {color: #b00000; font-weight: bold}
.note {font-style: italic; color: #444}
"""
CANDIDATE_FOOTER_H = 18


def _c_item(it: dict, mark: str = "") -> str:
    """A candidate A3 line: the mark (ENTERS / CHANGES), the id, the requirement, the class and the quoted consequence
    (never compacted: the candidate may run to several pages), source and flags."""
    parts = ([f'<span class="mark">[{_esc(mark)}]</span>'] if mark else []) + [f'<b>{_esc(it["id"])}</b> {_rich(it["text"])}']
    if it.get("consequence"):
        cls = f"<i>{_esc(it.get('class', ''))}</i>: " if it.get("class") else ""
        parts.append(f"— {cls}“{_rich(it['consequence'])}”")
        if it.get("gloss"):                              # session 11: labelled from the approval state (stage2)
            parts.append(f"({_esc(it.get('gloss_label') or 'proposed translation, not reviewed')}: ‘{_esc(it['gloss'])}’)")
    if it.get("source"):
        parts.append(f'<span class="meta">[{_rich(it["source"])}]</span>')
    if it.get("flags"):
        parts.append(f'<span class="flag">{SEP.join(_rich(f) for f in it["flags"])}</span>')
    return f'<p class="it">{" ".join(parts)}</p>'


def _c_body(c: dict) -> str:
    """The candidate A3 as HTML (tenderpack.partial.compute output): what may change, the changes against the validated
    A3 with their ops, every A3 section in full, the blockers, the conditional scenarios, the image-read units and the
    open issues. Presentation only."""
    a3 = c["a3"]
    stages = ", ".join(x["stage"] for x in c["pending"])
    pages = c.get("pages")
    out = [f'<div class="title">A3 CANDIDATE — what would put this bid out if the valid ops of {_esc(stages)} stood '
           f'(NOT VALIDATED)</div>',
           f'<div class="sub">{_rich(a3["subtitle"])}</div>',
           f'<div class="banner">{_esc(c["banner"])}. The validated A3 (a3.pdf, one page) stays at '
           f'{_esc(c["validated_stage"])} and is unchanged. This candidate is a view for review, not a deliverable: every '
           f'op in it is a proposal and nothing is accepted. It is not held to the one-page rule (C43 applies to the '
           f'validated A3 only)' + (f' and runs to {pages} page{"s" if pages != 1 else ""}.' if pages else ".") + '</div>',
           "<h2>What may be changing</h2>", f"<p>{_rich(c['paragraph'])}</p>",
           f"<h2>Changes against the validated A3 ({len(c['changes'])})</h2>"]
    for x in c["changes"]:
        out.append(f'<p class="it"><span class="mark">{_esc(x["change"].upper())}</span> <b>{_esc(x["row"])}</b> '
                   f'{_rich("; ".join(x["what"]))} <span class="meta">(row {_esc(x["status_validated"])} -&gt; '
                   f'{_esc(x["status_candidate"])}; review {_esc(x["row_review"])}' + ("; STALE" if x["stale"] else "")
                   + ")</span></p>")
        out += [f'<p class="sub2">source op {_rich(o["label"])}</p>' for o in x["ops"]]
        out += [f'<p class="sub2">{_rich(y)}</p>' for y in x["why"]]
    if not c["changes"]:
        out.append('<p class="note">No row enters, leaves or changes.</p>')
    marks = {x["row"]: x["change"].upper() for x in c["changes"]}
    out.append("<h2>The candidate A3, every section in full</h2>")
    for sec in a3["sections"]:
        out.append(f"<h3>{_rich(sec['heading'])}</h3>")
        if sec.get("note"):
            out.append(f'<p class="note">{_rich(sec["note"])}</p>')
        out += [_c_item(it, marks.get(it["id"], "")) for it in sec.get("items") or []]
        if sec.get("ids"):
            out.append('<p class="meta">' + ", ".join((f'<span class="mark">[{_esc(marks[i])}]</span> ' if i in marks else "")
                                                       + _esc(i) for i in sec["ids"]) + "</p>")
    b = c["blockers"]
    out.append("<h2>Blockers: what stops the candidate becoming the validated state, or what it cannot establish</h2>")
    out.append(f"<h3>Unresolved provisions ({len(b['provisions'])})</h3>")
    out += [f'<p class="it"><b>{_esc(x["provision"])}</b> ({_esc(x["stage"])} p{_esc(str(x["page"]))}; {_esc(x["kind"])}): '
            f'{_rich(x["reason"])} <span class="meta">— words: “{_rich(x["text"])}”; rows: '
            f'{_esc(", ".join(x["rows"]) or "none")}; activities: {_esc(", ".join(x["activities"]) or "none")}</span>'
            + (f' <span class="flag">— {_rich(x["route"])}</span>' if x.get("route") else "") + '</p>'
            for x in b["provisions"]] or ['<p class="note">None.</p>']
    if b["invalid_ops"] or b["withheld_ops"]:
        out.append(f"<h3>Invalid or withheld ops ({len(b['invalid_ops']) + len(b['withheld_ops'])})</h3>")
        out += [f'<p class="it">INVALID <b>{_esc(x["op"])}</b> ({_esc(x["provision"])}): {_rich(x["failed"])}</p>'
                for x in b["invalid_ops"]]
        out += [f'<p class="it">WITHHELD <b>{_esc(x["op"])}</b> ({_esc(x["provision"])}): rejected by a person</p>'
                for x in b["withheld_ops"]]
    out.append(f"<h3>Rows not settled in the candidate ({len(b['unresolved_rows'])})</h3>")
    out += [f'<p class="it"><b>{_esc(x["row"])}</b>: {_rich("; ".join(x["why"]))}</p>' for x in b["unresolved_rows"]] \
        or ['<p class="note">None.</p>']
    out.append(f"<h3>A5 activities blocked by an unresolved row ({len(b['blocked_activities'])})</h3>")
    out += [f'<p class="it"><b>{_esc(x["activity"])}</b> {_rich(x["name"])} <span class="meta">(rows '
            f'{_esc(", ".join(x["rows"]))})</span></p>' for x in b["blocked_activities"]] or ['<p class="note">None.</p>']
    out.append(f"<h3>Obligations reaching no output (C46) ({len(b['c46'])})</h3>")
    out += [f'<p class="it"><b>{_esc(x["op"])}</b> [{_esc(x["output"])}]: {_rich(x["detail"])}</p>' for x in b["c46"]] \
        or ['<p class="note">None.</p>']
    ch = b.get("chains") or []
    out.append(f"<h3>Relationship chains blocked or incomplete ({len(ch)})</h3>")
    out += [f'<p class="it"><b>{_esc(x["target"])}</b> via {_esc(" > ".join(x["path"]))}: '
            + "; ".join([f"BLOCKED: {_rich(y['document'] or '')} not supplied ({_esc(y['entry_id'] or '')}): cannot be "
                         f"established: {_rich(y['blocks'] or '')}" for y in x["blockers"]]
                        + ([f"CHAIN INCOMPLETE ({_esc(x['truncated_by'] or 'bound')}); not followed: "
                            f"{_esc(', '.join(x['unfollowed']))}"] if x["completeness"] == "truncated" else [])
                        + ([f"CYCLIC: {_esc(' > '.join(x['cycle']))}"] if x["completeness"] == "cyclic" else []))
            + "</p>" for x in ch] or ['<p class="note">None.</p>']
    out.append(f"<h3>Documents referenced but not supplied ({len(b['missing_documents'])}): kept visible</h3>")
    out += [f'<p class="it"><b>{_esc(x["id"])}</b> ' + (f"{_rich(x['document'])}: " if x["document"] else "")
            + _rich(x["blocks"]) + ' <span class="meta">— ' + (
                f"reached by this addendum: {_esc(', '.join(x['reached_by_this_addendum'][:8]))}"
                if x["reached_by_this_addendum"] else "not reached by this addendum's changes; still open")
            + (f"; issues {_esc(', '.join(x['issues']))}" if x["issues"] else "") + "</span></p>"
            for x in b["missing_documents"]] or ['<p class="note">None.</p>']
    out.append(f"<h3>Conflicts ({len(b['conflicts'])})</h3>")
    out += [f'<p class="it">{_rich(x["what"])} <span class="meta">({_esc(x["source"])})</span></p>' for x in b["conflicts"]] \
        or ['<p class="note">None.</p>']
    if b["op_issues"]:
        out.append(f"<h3>Issues the ops raise ({len(b['op_issues'])})</h3>")
        out += [f'<p class="it"><b>{_esc(x["op"])}</b>{"" if x["applied"] else " (not applied)"}: {_rich(x["issue"])}</p>'
                for x in b["op_issues"]]
    out.append(f"<h3>Open issues in play: kept open, never resolved here ({len(c['open_issues_in_play'])})</h3>")
    out += [f'<p class="it"><b>{_esc(x["id"])}</b> ({_esc(x["owner"])}): {_rich(x["text"])} <span class="meta">— via '
            f'{_esc(", ".join(x["via"][:8]))}</span></p>' for x in c["open_issues_in_play"]] or ['<p class="note">None.</p>']
    out.append(f"<h2>Conditional scenarios ({len(c['scenarios'])})</h2>")
    for x in c["scenarios"]:
        out.append(f'<p class="it"><b>{_esc(x["provision"])}</b> ({_esc(x["kind"])} {_esc(x.get("what") or "")}; '
                   f'{_esc(x["source"])}; '
                   f'{"applied" if x["applied"] else "not applied"} in the candidate)'
                   + (f": decision by <b>{_esc(x['trigger_deadline'])}</b>" if x["trigger_deadline"] else "")
                   + (f"; effective from {_esc(x['effective_from'])}" if x["effective_from"] else "") + "</p>")
        out += [f'<p class="sub2">trigger words: “{_rich(x["trigger"])}”</p>' if x["trigger"] else "",
                f'<p class="sub2">if triggered: {_rich(x["if_triggered"])}; if not: {_rich(x["if_not_triggered"])}</p>',
                f'<p class="sub2">rows: {_esc(", ".join(x["rows"][:10]) or "none")}; {_esc(x["model"])}</p>']
    if not c["scenarios"]:
        out.append('<p class="note">None in the op file: no op or unresolved provision says it is conditional or '
                   'effective-dated.</p>')
    out.append("<h2>Image-read units: crops, Arabic verbatim, translations apart, table context, units, uncertainties</h2>")
    out.append('<p class="note">The review packet of each region (in the evidence build) shows every crop beside its '
               'reading. Translations are proposals, never evidence; what a value means (maximum, range, target) is '
               'decided by a person.</p>')
    for g in c["images"]:
        out.append(f'<p class="it"><b>{_esc(g["region"])}</b> ({_esc(g["doc"])} p{_esc(str(g["page"]))}; reading '
                   f'{_esc(g["status"])}): packet &lt;evidence build&gt;/{_esc(g["packet"])}'
                   + ("" if g["touched"] else f" — not touched by {_esc(c['stage'])}; kept available") + "</p>")
        out += [f'<p class="sub2">uncertainty: {_rich(u)}</p>' for u in g["uncertainties"][:8]]
        for u in g["units"]:
            words = u["text"] or "; ".join(f"{k}: {v}" for k, v in u["cells"].items())
            bits = [f"<b>{_esc(u['unit'])}</b> “{_rich(words)}”"]
            if u["translation"]:
                bits.append(f"translation (a proposal, not evidence): ‘{_rich(u['translation'])}’")
            if u["measure"]:
                bits.append(f"unit {_esc(u['measure'])}")
            if u["table"].get("column_headings"):
                bits.append(f"table: {_rich(u['table'].get('table_title') or '')} columns "
                            f"{_esc(', '.join(u['table']['column_headings']))}"
                            + (f"; qualifier {_rich(u['table']['qualifier'])}" if u["table"].get("qualifier") else ""))
            if u["candidate"]:
                cand = u["candidate"]
                bits.append(f"at {_esc(c['stage'])}: {_esc(cand['status'])} “{_rich(cand['text'] or '; '.join(f'{k}: {v}' for k, v in (cand['cells'] or {}).items()))}”"
                            f" (by {_esc(', '.join(cand['changed_by']))})")
            if u["uncertain"]:
                bits.append(f"uncertain: {_rich('; '.join(u['uncertain']))}")
            if u["issues"]:
                bits.append(f"issues {_esc(', '.join(u['issues']))}")
            bits.append(f'<span class="meta">touched: {_rich("; ".join(u["touched"]))}</span>')
            out.append(f'<p class="sub2">{SEP.join(bits)}</p>')
    g = a3.get("groups") or {}
    out.append(f"<h2>{_rich(g.get('heading', 'Open issues'))}</h2>")
    for grp in g.get("groups") or []:
        out.append(f'<p class="it"><b>{_rich(grp["title"])}</b> ({len(grp["items"])}): '
                   + "; ".join(f"<b>{_esc(i['id'])}</b> {_rich(i['short'])}" for i in grp["items"])
                   + (f' <span class="meta">questions drafted: {_esc(", ".join(grp["questions"]))}</span>'
                      if grp.get("questions") else "") + "</p>")
    out.append(f"<h2>Ops standing in the candidate ({len(c['ops'])})</h2>")
    out += [f'<p class="it">{_rich(o["label"])}</p>' for o in c["ops"]] or ['<p class="note">None.</p>']
    return "\n".join(x for x in out if x)


def candidate_a3_html(c: dict, standalone: bool = False) -> str:
    """The candidate A3 as HTML: the body for the PDF, or (standalone) a page for a browser."""
    body = _c_body(c)
    if not standalone:
        return body
    return ("<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" "
            "content=\"width=device-width, initial-scale=1\"><title>A3 candidate</title><style>"
            "body{font-family:system-ui,sans-serif;margin:16px;max-width:1100px;color:#111;background:#fff;font-size:14px;"
            "line-height:1.35}.title{font-size:20px;font-weight:bold}.banner{background:#fde2e2;border:2px solid #b00000;"
            "color:#7a0000;font-weight:bold;padding:6px 10px;margin:8px 0}.it{padding-left:14px;text-indent:-14px;"
            "margin:2px 0}.sub2{padding-left:28px;color:#333;margin:1px 0}.meta{color:#555}.flag,.mark{color:#b00000}"
            ".mark{font-weight:bold}.note{font-style:italic;color:#555}h2{border-bottom:1px solid #999}"
            "</style></head><body>" + body + "</body></html>\n")


def write_candidate_a3_pdf(c: dict, path: Path) -> dict:
    """The candidate A3 on as many A4 pages as it needs (pymupdf Story), each with a footer 'A3 CANDIDATE (not
    validated) ... page i of N'; the banner states N (laid out twice when the count changes). Deterministic: fixed
    metadata, no /ID, no links. Returns {"pages": N}."""
    w, h = A3_PAGE.width, A3_PAGE.height
    where = pymupdf.Rect(A3_MARGIN_X, A3_MARGIN_Y, w - A3_MARGIN_X, h - A3_MARGIN_Y - CANDIDATE_FOOTER_H - 4)

    def layout(n):
        story = pymupdf.Story(html=_c_body(dict(c, pages=n)), user_css=CANDIDATE_CSS)
        buf = io.BytesIO()
        wr = pymupdf.DocumentWriter(buf)
        more = 1
        while more:
            dev = wr.begin_page(A3_PAGE)
            more, _ = story.place(where)
            story.draw(dev)
            wr.end_page()
        wr.close()
        return pymupdf.open("pdf", buf.getvalue())
    n, src = None, None
    for _ in range(3):
        src = layout(n)
        if src.page_count == n:
            break
        n = src.page_count
    doc = pymupdf.open()
    doc.insert_pdf(src)
    stage = c["stage"]
    for i, page in enumerate(doc, 1):
        foot = pymupdf.Rect(A3_MARGIN_X, h - A3_MARGIN_Y - CANDIDATE_FOOTER_H, w - A3_MARGIN_X, h - A3_MARGIN_Y)
        page.insert_htmlbox(foot, f'<div class="foot">A3 CANDIDATE (not validated) — {_esc(stage)} as proposed; the '
                                  f'validated A3 is a3.pdf — page {i} of {doc.page_count}</div>',
                            css=A3_FOOTER_CSS, scale_low=A3_SCALE_LOW)
    doc.subset_fonts()
    doc.set_metadata({"title": f"A3 CANDIDATE {stage} (not validated)", "creator": PRODUCER, "producer": PRODUCER,
                      "creationDate": PDF_DATE, "modDate": PDF_DATE})
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(doc.tobytes(garbage=3, deflate=True, no_new_id=True))
    pages = doc.page_count
    doc.close()
    return {"pages": pages}
